import os
import shutil
import fitz


def _is_scanned_pdf(text: str) -> bool:
    return len(text.strip()) < 150


def extract_text_from_pdf(pdf_path: str) -> str:
    document = fitz.open(pdf_path)
    text = ""
    for page in document:
        text += page.get_text()
    document.close()
    return text


def extract_text_from_pdf_ocr(pdf_path: str) -> str:
    try:
        import pytesseract
        from PIL import Image
        import io

        tesseract_cmd = os.getenv("TESSERACT_CMD") or shutil.which("tesseract")
        if not tesseract_cmd:
            return "[Scanned PDF detected - Tesseract is not installed or is not on PATH. Install it with 'brew install tesseract'.]"
        pytesseract.pytesseract.tesseract_cmd = tesseract_cmd

        document = fitz.open(pdf_path)
        text = ""
        for page_num, page in enumerate(document):
            mat = fitz.Matrix(300 / 72, 300 / 72)
            pix = page.get_pixmap(matrix=mat)
            img_bytes = pix.tobytes("png")
            image = Image.open(io.BytesIO(img_bytes))
            page_text = pytesseract.image_to_string(image, lang="eng")
            text += f"\n[Page {page_num + 1}]\n{page_text}"
        document.close()
        return text

    except ImportError:
        return "[Scanned PDF detected — OCR not available. Install pytesseract and Pillow.]"

    except Exception as e:
        return f"[OCR failed: {str(e)}]"


def extract_text_from_docx(docx_path: str) -> str:
    try:
        from docx import Document
        doc = Document(docx_path)
        paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]

        table_texts = []
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    table_texts.append(row_text)

        all_text = "\n".join(paragraphs)
        if table_texts:
            all_text += "\n\n[Tables]\n" + "\n".join(table_texts)

        return all_text

    except ImportError:
        return "[DOCX extraction failed — install python-docx.]"

    except Exception as e:
        return f"[DOCX extraction failed: {str(e)}]"


def extract_text(file_path: str) -> str:
    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".docx":
        return extract_text_from_docx(file_path)

    elif ext == ".pdf":
        text = extract_text_from_pdf(file_path)
        if _is_scanned_pdf(text):
            return extract_text_from_pdf_ocr(file_path)
        return text

    else:
        return f"[Unsupported file type: {ext}. Supported: .pdf, .docx]"